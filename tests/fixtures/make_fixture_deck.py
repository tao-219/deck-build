#!/usr/bin/env python3
"""make_fixture_deck.py — build a tiny NEUTRAL reference .pptx from a plan fixture.

Renders one slide per plan slide using the deck-build Path-B helpers
(`deck_helpers.two_tone_title`), so the smoke test has a real .pptx to lint /
--semantic without committing a binary. Synthetic, brand-neutral content only
(no client data) — palette + fonts overridden to generic values so the fixture
carries no brand fingerprint.

The title shape is named "Title 1" so `action_title_lint.extract_title` resolves
it via its priority-1 path (same as `deck_helpers.replace_title`).

Usage (standalone):
  python make_fixture_deck.py <plan.json> <out.pptx>

Importable:
  from make_fixture_deck import build_deck
  build_deck(plan_dict, "/tmp/good.pptx")
"""

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

# Path-B helpers from the plugin's importable module.
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
import deck_helpers as dh  # noqa: E402


def build_deck(plan: dict, out_path):
    """Build a minimal neutral deck: one slide per plan slide, two-tone action title."""
    # Brand-neutral DNA so the fixture carries no client/brand fingerprint.
    dh.set_fonts(heading="Arial", body="Arial")
    dh.set_palette(primary=RGBColor(0x1F, 0x4E, 0x79),          # neutral navy
                   text_dark=RGBColor(0x20, 0x20, 0x20))

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]  # blank layout

    for s in plan.get("slides", []):
        slide = prs.slides.add_slide(blank)
        bold = s.get("title_bold") or ""
        normal = s.get("title_normal")
        box = dh.two_tone_title(slide, 0.5, 0.55, 12.3, 0.9, bold, normal, size_pt=22)
        # Name it "Title 1" so the linter's title resolver finds it deterministically.
        box.name = "Title 1"

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return str(out_path)


def main():
    if len(sys.argv) != 3:
        print("usage: python make_fixture_deck.py <plan.json> <out.pptx>", file=sys.stderr)
        sys.exit(2)
    plan = json.load(open(sys.argv[1]))
    path = build_deck(plan, sys.argv[2])
    print(f"wrote {path} ({len(plan.get('slides', []))} slides)")


if __name__ == "__main__":
    main()
