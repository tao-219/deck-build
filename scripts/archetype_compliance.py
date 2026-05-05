#!/usr/bin/env python3
"""archetype_compliance.py — XML-walking QA for deck-build outputs.

Catches the high-leverage defects without needing a vision LLM call:

  1. Footer-clearance overrun — content shape extends past the y-coord limit
     for its archetype (per design-archetypes.md "Footer-clearance budget").
  2. Container clipping — a child shape extends below its container shape
     (specifically "container" = a rectangle/rounded-rectangle with dashed
     border line).
  3. Two-tone title check — title with >8 words but only single-weight runs.
  4. Title-as-noun-phrase — title lacks a verb (heuristic) — POTENTIAL
     action-title violation.

Usage:
  python archetype_compliance.py <deck.pptx> [--plan plan.json]

If --plan is provided, expected archetypes are read per slide; otherwise the
default footer-clearance limit (6.65") is used for every slide.

Returns JSON to stdout with structure:
  {
    "deck": "<path>",
    "slide_count": N,
    "issues": [
      {"slide": N, "severity": "High|Medium|Low",
       "rule": "footer-clearance|container-clipping|two-tone-title|verb-test",
       "shape": "<name>",
       "description": "..."}
    ]
  }
Exit code: 0 if no High issues, 1 otherwise.
"""

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

# Share title-extraction + verb/two-tone logic with the action_title_lint module
sys.path.insert(0, str(Path(__file__).parent))
from action_title_lint import extract_title, has_verb, has_two_tone, propose_two_tone_split

# Default footer clearance — same map as deck_helpers.FOOTER_CLEARANCE
FOOTER_CLEARANCE = {
    "stage-flow-with-signoff":          6.65,
    "two-col-overview-with-subcallout": 6.65,
    "spotlight-comparison-columns":     6.85,
    "tiered-region-coverage-table":     6.55,
    "recommendation-hero":              6.65,
    "before-after-process":             6.55,
    "risk-heatmap":                     6.55,
    "default":                          6.65,
}

# Names that ARE the footer band — exempt from footer-clearance check
FOOTER_SHAPE_PATTERNS = {
    "Slide Number Placeholder",
    "Copyright",
    "Footer",
    "Manulife",  # Manulife logo if it carries that name
}

# Verb/two-tone logic now lives in action_title_lint (imported above).


def is_footer_shape(name: str) -> bool:
    return any(pat in name for pat in FOOTER_SHAPE_PATTERNS)


def emu_to_inches(value):
    if value is None:
        return None
    return value / 914400.0


def check_two_tone(slide) -> tuple[bool, str | None]:
    """Return (passes_check, title_text). Passes if title is ≤8 words OR has mixed weight."""
    text, paragraph = extract_title(slide)
    if text is None:
        return True, None
    word_count = len(text.split())
    if word_count <= 8:
        return True, text
    return has_two_tone(paragraph), text


def check_verb(text: str) -> bool:
    """Heuristic: title is OK if it passes the action_title_lint verb test."""
    if not text:
        return True
    return has_verb(text)


def get_shape_bounds(shape) -> tuple[float, float, float, float] | None:
    """Return (x, y, w, h) in inches, or None if any coord is missing."""
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return (emu_to_inches(shape.left),
            emu_to_inches(shape.top),
            emu_to_inches(shape.width),
            emu_to_inches(shape.height))


def is_dashed_border(shape) -> bool:
    """Check if shape has a dashed-border line element."""
    try:
        # Walk the shape XML for prstDash
        sp = shape._element
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        dash_elements = sp.findall(".//a:prstDash", ns)
        return len(dash_elements) > 0
    except Exception:
        return False


def is_container_shape(shape) -> bool:
    """Container = AUTO_SHAPE rectangle/rounded-rectangle with dashed border, ≥4" wide.

    Tables, groups, and placeholders are NOT containers even if they have
    dashed cell borders.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
        return False
    bounds = get_shape_bounds(shape)
    if bounds is None:
        return False
    x, y, w, h = bounds
    if w < 4.0:
        return False
    return is_dashed_border(shape)


def check_slide(slide, slide_num, archetype="default"):
    issues = []
    clearance = FOOTER_CLEARANCE.get(archetype, FOOTER_CLEARANCE["default"])

    # Find containers
    containers = []
    for shape in slide.shapes:
        if is_container_shape(shape):
            bounds = get_shape_bounds(shape)
            if bounds:
                containers.append((shape.name, bounds))

    # Walk all non-footer shapes
    for shape in slide.shapes:
        if is_footer_shape(shape.name):
            continue
        bounds = get_shape_bounds(shape)
        if bounds is None:
            continue
        x, y, w, h = bounds
        bottom = y + h

        # Footer-clearance check
        if bottom > clearance + 0.05:  # 0.05" tolerance
            delta = bottom - clearance
            # Suggest concrete remediation: shrink height or row count
            fix_hint_parts = [
                f"reduce shape height by {delta:.2f}\" (current {h:.2f}\" → {max(0.1, h - delta):.2f}\")"
            ]
            # If shape is tall enough that it likely contains rows, estimate row reduction
            if h > 0.5:
                # Heuristic: typical row height 0.20–0.30"
                est_row_h = 0.25
                rows_to_drop = max(1, int((delta / est_row_h) + 0.5))
                fix_hint_parts.append(
                    f"OR drop ~{rows_to_drop} row(s) at est. row-height 0.25\""
                )
            # If shape origin is high on the slide, raising start y won't help — note it
            if y < 1.5:
                fix_hint_parts.append(
                    f"shape starts at y={y:.2f}\" (top of content area); shrinking is the only option"
                )
            issues.append({
                "slide": slide_num,
                "severity": "High",
                "rule": "footer-clearance",
                "shape": shape.name,
                "description": (
                    f"Bottom edge at y={bottom:.2f}\" exceeds clearance limit "
                    f"y={clearance:.2f}\" for archetype '{archetype}' (over by {delta:.2f}\")."
                ),
                "suggested_fix": " — ".join(fix_hint_parts),
            })

        # Container clipping check — does this shape extend below any container
        # whose x-range it overlaps?
        if not is_container_shape(shape):
            for c_name, (cx, cy, cw, ch) in containers:
                # Overlap in x?
                if x + w < cx or x > cx + cw:
                    continue
                # Is shape mostly INSIDE the container's x range?
                if x < cx - 0.1 or x + w > cx + cw + 0.1:
                    continue
                # Then it should be bounded by the container's y extent
                if bottom > cy + ch + 0.05:
                    delta = bottom - (cy + ch)
                    issues.append({
                        "slide": slide_num,
                        "severity": "High",
                        "rule": "container-clipping",
                        "shape": shape.name,
                        "description": (
                            f"Shape bottom y={bottom:.2f}\" extends below "
                            f"container '{c_name}' bottom y={cy + ch:.2f}\" "
                            f"(over by {delta:.2f}\")."
                        ),
                        "suggested_fix": (
                            f"extend container '{c_name}' height by {delta:.2f}\" "
                            f"(current {ch:.2f}\" → {ch + delta:.2f}\"), "
                            f"OR shrink child shape '{shape.name}' height by {delta:.2f}\""
                        ),
                    })

    # Two-tone title check
    passes, title = check_two_tone(slide)
    if not passes and title:
        # Propose a split point
        split = propose_two_tone_split(title)
        suggested_fix = (
            f"Split as bold='{split['bold']}' + normal='{split['normal']}' "
            f"(split at {split['rationale']})"
        ) if split else "Split into a bold short claim + normal-weight elaboration"
        issues.append({
            "slide": slide_num,
            "severity": "Medium",
            "rule": "two-tone-title",
            "shape": "Title",
            "description": (
                f"Title is {len(title.split())} words but rendered as a single "
                f"weight throughout. Apply the two-tone pattern (bold short "
                f"claim + en-dash + normal-weight elaboration). Title: \"{title[:80]}\""
            ),
            "suggested_fix": suggested_fix,
        })

    # Verb check
    if title and not check_verb(title):
        issues.append({
            "slide": slide_num,
            "severity": "Medium",
            "rule": "verb-test",
            "shape": "Title",
            "description": (
                f"Title appears to be a noun phrase (no verb detected by "
                f"heuristic). Action titles state a conclusion. "
                f"Title: \"{title[:80]}\""
            ),
        })

    return issues


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("deck", help="Path to .pptx to audit")
    parser.add_argument("--plan", help="Optional plan.json mapping slide_num → archetype")
    parser.add_argument("--quiet", action="store_true", help="Only print JSON, not human summary")
    args = parser.parse_args()

    deck_path = Path(args.deck)
    if not deck_path.exists():
        print(f"ERROR: {deck_path} does not exist", file=sys.stderr)
        sys.exit(2)

    plan = {}
    if args.plan:
        plan_path = Path(args.plan)
        if plan_path.exists():
            with open(plan_path) as f:
                plan_data = json.load(f)
            # Expect plan_data to have a "slides" array with per-slide archetype
            for slide_spec in plan_data.get("slides", []):
                idx = slide_spec.get("slide_number") or slide_spec.get("index")
                if idx is not None:
                    plan[int(idx)] = slide_spec.get("archetype", "default")

    prs = Presentation(str(deck_path))
    all_issues = []
    for i, slide in enumerate(prs.slides, 1):
        archetype = plan.get(i, "default")
        issues = check_slide(slide, i, archetype)
        all_issues.extend(issues)

    result = {
        "deck": str(deck_path),
        "slide_count": len(prs.slides),
        "issues": all_issues,
    }
    print(json.dumps(result, indent=2))

    if not args.quiet:
        high_count = sum(1 for issue in all_issues if issue["severity"] == "High")
        med_count  = sum(1 for issue in all_issues if issue["severity"] == "Medium")
        print(f"\n--- Summary: {len(all_issues)} issues "
              f"({high_count} High, {med_count} Medium) ---", file=sys.stderr)

    sys.exit(1 if any(i["severity"] == "High" for i in all_issues) else 0)


if __name__ == "__main__":
    main()
